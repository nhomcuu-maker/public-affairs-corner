"""Goc Dan Sinh: allowlisted document intake, plain-text extract and legal chunk index.
Python 3.10+, pip install -r requirements.txt. No site discovery/bypass.
"""
import argparse,csv,datetime,hashlib,json,re,time,unicodedata,io,os,zipfile
from pathlib import Path
from urllib.parse import urlparse
from bs4 import BeautifulSoup
import fitz
from docx import Document

ROOT=Path(__file__).resolve().parents[2]
DATA=ROOT/'data'/'law_rag'
RAW=DATA/'raw'; TEXT=DATA/'text'; PROCESSED=DATA/'processed'
ALLOWED={'vbpl.vn','vanban.chinhphu.vn','congbao.chinhphu.vn','dichvucong.gov.vn','thutuc.dichvucong.gov.vn'}
DOMAINS={'LEGAL','ADMIN'}
LOCAL_DOMAINS={'LEGAL','ADMIN','SCAM'}
LOCAL_EXTENSIONS={'.pdf','.docx','.txt','.html','.htm','.json','.jsonl','.csv'}
MAX_LOCAL_BYTES=25_000_000
LEGAL={'lao-dong','dan-su','dat-dai','hon-nhan'}
ADMIN={'con-nho','hoc-tap','viec-lam','cu-tru','hon-nhan','dat-dai','y-te','phuong-tien','huu-tri','tu-vong','khieu-kien'}

def host_ok(url):
 p=urlparse(url)
 return p.scheme=='https' and (p.hostname in ALLOWED or any(p.hostname.endswith('.'+d) for d in ALLOWED)) and not p.username and not p.password

def normalize(t):return unicodedata.normalize('NFC',str(t or '')).strip()
def ensure():
 for p in (RAW,TEXT,PROCESSED):p.mkdir(parents=True,exist_ok=True)
def read_seeds():
 p=DATA/'seeds.csv'
 if not p.exists():raise SystemExit('No seeds.csv; add verified document URLs to data/law_rag/seeds.csv')
 for row in csv.DictReader(p.open(encoding='utf-8-sig',newline='')):
  url=(row.get('url') or '').strip();domain=(row.get('domain') or '').strip().upper();topic=(row.get('topic') or '').strip();title=normalize(row.get('title'))
  if not url:continue
  if not host_ok(url):print('SKIP: non-allowlisted URL',url);continue
  if domain not in DOMAINS or topic not in (LEGAL if domain=='LEGAL' else ADMIN):print('SKIP: invalid taxonomy',url);continue
  if not title or len(title)<12:print('SKIP: missing meaningful Vietnamese title',url);continue
  yield {'url':url,'domain':domain,'topic':topic,'title':title,'doc_number':normalize(row.get('doc_number')),'effective_from':normalize(row.get('effective_from'))}

def extract(b,ctype,url):
 suffix=Path(urlparse(url).path).suffix.lower()
 if b[:4]==b'%PDF' or suffix=='.pdf':
  with fitz.open(stream=b,filetype='pdf') as pdf:return '\n'.join(p.get_text('text') for p in pdf)
 if suffix=='.docx' or 'wordprocessingml' in ctype:
  import io
  return '\n'.join(p.text for p in Document(io.BytesIO(b)).paragraphs)
 if 'html' in ctype or suffix in ('','.html','.htm'):
  soup=BeautifulSoup(b,'html.parser')
  for el in soup(['script','style','nav','footer','header','form']):el.decompose()
  target=soup.find('article') or soup.find('main') or soup.body or soup
  return target.get_text('\n',strip=True)
 return ''

def collect():
 import requests  # network dependency is not required by offline local-file intake
 ensure();records=[]
 with requests.Session() as ses:
  ses.headers.update({'User-Agent':'GocDanSinh-Research/1.0 (respectful document intake)'})
  for seed in read_seeds():
   u=seed['url'];ident=hashlib.sha256(u.encode()).hexdigest()[:20]
   try:
    r=ses.get(u,timeout=30,allow_redirects=False)
    if r.status_code!=200:raise ValueError('HTTP '+str(r.status_code))
    if len(r.content)>25_000_000:raise ValueError('document exceeds 25MB')
    ctype=r.headers.get('Content-Type','').split(';')[0].lower()
    if ctype not in ('application/pdf','text/html','application/xhtml+xml','application/vnd.openxmlformats-officedocument.wordprocessingml.document','application/octet-stream'):raise ValueError('unsupported MIME '+ctype)
    if ctype=='application/octet-stream' and not r.content.startswith(b'%PDF'):raise ValueError('unknown binary')
    text=normalize(extract(r.content,ctype,u))
    if len(text)<120:raise ValueError('too little extractable text; scanned or page shell')
    suffix=('.pdf' if r.content.startswith(b'%PDF') else '.docx' if '.docx' in u.lower() else '.html')
    (RAW/(ident+suffix)).write_bytes(r.content);(TEXT/(ident+'.txt')).write_text(text,encoding='utf-8')
    records.append({**seed,'id':ident,'source_url':u,'content_sha256':hashlib.sha256(r.content).hexdigest(),'captured_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'raw_file':str(Path('data/law_rag/raw')/(ident+suffix)),'text_file':str(Path('data/law_rag/text')/(ident+'.txt')),'text_chars':len(text),'verified':False})
    print('OK',ident,seed['topic'],len(text),'chars')
   except Exception as e:print('FAILED',u,str(e))
   time.sleep(1.2)
 out=PROCESSED/'documents.jsonl';old={}
 if out.exists():
  for line in out.read_text(encoding='utf-8').splitlines():
   if line.strip():
    x=json.loads(line);old[x['id']]=x
 for x in records:old[x['id']]=x
 out.write_text(''.join(json.dumps(x,ensure_ascii=False)+'\n' for x in old.values()),encoding='utf-8')
 print('COLLECTED',len(records),'INDEXED METADATA TOTAL',len(old))

def chunk_text(text,maxlen=1250):
 # Preserve article headings as anchors instead of blind fixed-size cuts.
 parts=re.split(r'(?=\n\s*(?:Điều\s+\d+[.:]|Chương\s+[IVX]+\b|Mục\s+\d+\b))',text,flags=re.I)
 buf=''
 for part in parts:
  for paragraph in re.split(r'\n\s*\n',part):
   paragraph=paragraph.strip()
   if not paragraph:continue
   if len(buf)+len(paragraph)+1>maxlen and buf:
    yield buf;buf=''
   while len(paragraph)>maxlen:
    # Prefer a sentence boundary, but ensure progress.
    cut=paragraph.rfind('. ',0,maxlen)
    if cut<maxlen//2:cut=maxlen
    else:cut+=1
    yield paragraph[:cut].strip();paragraph=paragraph[cut:].strip()
   buf=(buf+'\n'+paragraph).strip()
 if buf:yield buf



def extract_local_content(filename, payload):
 """Untrusted attachments become plain text only; never interpreted as instructions."""
 extension=Path(filename).suffix.lower()
 if extension=='.pdf':
  if not payload.startswith(b'%PDF-'):raise ValueError('PDF signature mismatch')
  with fitz.open(stream=payload,filetype='pdf') as pdf:
   if len(pdf)>1000:raise ValueError('PDF exceeds page limit')
   return '\n'.join(page.get_text('text') for page in pdf)
 if extension=='.docx':
  if not payload.startswith(b'PK') or not zipfile.is_zipfile(io.BytesIO(payload)):
   raise ValueError('DOCX signature mismatch')
  with zipfile.ZipFile(io.BytesIO(payload)) as z:
   names=z.namelist()
   if len(names)>2500 or sum(i.file_size for i in z.infolist())>100_000_000:
    raise ValueError('DOCX expanded size too large')
   if 'word/document.xml' not in names:raise ValueError('Not a Word DOCX')
  doc=Document(io.BytesIO(payload))
  lines=[p.text for p in doc.paragraphs]
  for table in doc.tables:
   for row in table.rows:lines.append(' | '.join(cell.text for cell in row.cells))
  return '\n'.join(lines)
 if extension not in LOCAL_EXTENSIONS:raise ValueError('Unsupported extension')
 plain=payload.decode('utf-8-sig')
 if extension in {'.html','.htm'}:
  soup=BeautifulSoup(plain,'html.parser')
  for el in soup(['script','style','noscript','nav','footer','header','form']):el.decompose()
  return soup.get_text('\n',strip=True)
 if extension=='.json':
  return json.dumps(json.loads(plain),ensure_ascii=False,indent=2)
 if extension=='.jsonl':
  return '\n'.join(json.dumps(json.loads(line),ensure_ascii=False) for line in plain.splitlines() if line.strip())
 if extension=='.csv':
  rows=csv.reader(io.StringIO(plain))
  return '\n'.join(' | '.join(row) for row in rows)
 return plain


def import_local(input_dir, manifest_path):
 """Stage admin-provided files for HUMAN review. No SQL, public or approval writes."""
 folder=Path(input_dir).resolve(strict=True)
 if not folder.is_dir():raise ValueError('input-dir must be a directory')
 manifest=Path(manifest_path).resolve(strict=True)
 if not manifest.is_file():raise ValueError('manifest must be a CSV file')
 if manifest.parent!=folder:raise ValueError('manifest must be inside input-dir')
 ensure()
 index=PROCESSED/'documents.jsonl'
 records={}
 if index.exists():
  for line in index.read_text(encoding='utf-8').splitlines():
   if line.strip():
    obj=json.loads(line);records[obj['id']]=obj
 imported=0;unchanged=0;requires_ocr=0;failures=[]
 with manifest.open(encoding='utf-8-sig',newline='') as input_csv:
  reader=csv.DictReader(input_csv)
  if not reader.fieldnames or not {'filename','title','domain','topic'} <= set(reader.fieldnames):
   raise ValueError('manifest needs filename,title,domain,topic columns')
  for rowno, row in enumerate(reader,2):
   name=(row.get('filename') or '').strip()
   try:
    if not name or name!=Path(name).name or '/' in name or '\\' in name or name in {'.','..'}:
     raise ValueError('Filename must be a direct child, not a path')
    file=folder/name
    if file.is_symlink() or not file.is_file():raise ValueError('Missing file or symbolic link')
    ext=file.suffix.lower()
    if ext not in LOCAL_EXTENSIONS:raise ValueError('Unsupported extension')
    size=file.stat().st_size
    if size<1 or size>MAX_LOCAL_BYTES:raise ValueError('File empty or above 25 MB')
    domain=(row.get('domain') or '').strip().upper()
    title=normalize(row.get('title'))
    topic=normalize(row.get('topic'))
    if domain not in LOCAL_DOMAINS:raise ValueError('Unrecognized domain')
    if len(title)<8 or len(title)>400 or len(topic)<2 or len(topic)>100:
     raise ValueError('Missing/invalid title or topic')
    source_url=(row.get('source_url') or '').strip()
    if source_url and not host_ok(source_url):
     raise ValueError('Source URL must be an approved official HTTPS host')
    payload=file.read_bytes()
    digest=hashlib.sha256(payload).hexdigest()
    identity=digest[:24]
    if identity in records:
     if records[identity].get('content_sha256')!=digest:
      raise ValueError('SHA identity collision')
     # Repair legacy PDF records that contained only signer metadata/page markers.
     existing=records[identity]
     if (ext=='.pdf' and existing.get('origin')=='local_admin_upload'
       and existing.get('verified') is False and existing.get('intake_status')=='pending_review'
       and int(existing.get('text_chars') or 0)<350):
      existing['intake_status']='requires_ocr'
      requires_ocr+=1
     unchanged+=1;continue
    text=normalize(extract_local_content(name,payload))
    if len(text)>2_000_000:raise ValueError('Extracted text too large')
    needs_ocr=ext=='.pdf' and len(text)<350
    if len(text)<120 and not needs_ocr:raise ValueError('Too little extractable text')
    staged_name=identity+ext
    target=RAW/staged_name
    if target.exists() and hashlib.sha256(target.read_bytes()).hexdigest()!=digest:
     raise ValueError('Immutable file hash mismatch')
    if not target.exists():
     temp=RAW/(staged_name+'.tmp');temp.write_bytes(payload);os.replace(temp,target)
    text_name=identity+'.txt'
    if text:
     target_text=TEXT/text_name
     temp_text=TEXT/(text_name+'.tmp');temp_text.write_text(text,encoding='utf-8');os.replace(temp_text,target_text)
    now=datetime.datetime.now(datetime.timezone.utc).isoformat()
    records[identity]={
     'id':identity,'title':title,'domain':domain,'topic':topic,
     'url':source_url,'source_url':source_url,'doc_number':normalize(row.get('doc_number')),
     # Dates and legal validity are NOT inferred from uploaded text or filename.
     'effective_from':normalize(row.get('effective_from')),
     'content_sha256':digest,'captured_at':now,
     'raw_file':str(Path('data/law_rag/raw')/staged_name),
     'text_file':str(Path('data/law_rag/text')/text_name) if text else '',
     'text_chars':len(text),'verified':False,'published':False,
     'intake_status':'requires_ocr' if needs_ocr else 'pending_review',
     'origin':'local_admin_upload','filename':name
    }
    imported+=1
    if needs_ocr:requires_ocr+=1
   except (OSError,UnicodeError,ValueError,RuntimeError,zipfile.BadZipFile) as exc:
    failures.append({'row':rowno,'filename':name,'error':str(exc)[:180]})
 # No partially written JSONL file is ever visible to subsequent pipeline stages.
 temp=index.with_suffix('.jsonl.tmp')
 temp.write_text(''.join(json.dumps(x,ensure_ascii=False)+'\n' for x in records.values()),encoding='utf-8')
 os.replace(temp,index)
 report={'imported':imported,'unchanged':unchanged,'requires_ocr':requires_ocr,
         'failed':len(failures),'failures':failures,
         'verification':'PENDING_REVIEW','published':False}
 print(json.dumps(report,ensure_ascii=False,indent=2))
 return report

def process():
 ensure();p=PROCESSED/'documents.jsonl'
 if not p.exists():raise SystemExit('First run collect')
 chunks=[]
 for line in p.read_text(encoding='utf-8').splitlines():
  if not line.strip():continue
  doc=json.loads(line)
  if doc.get('intake_status')=='requires_ocr' or not doc.get('text_file'):continue
  file=ROOT/doc['text_file']
  if not file.is_file():continue
  for i,content in enumerate(chunk_text(file.read_text(encoding='utf-8'))):
   chunks.append({'id':doc['id']+'-'+str(i),'document_id':doc['id'],'title':doc['title'],'domain':doc['domain'],'topic':doc['topic'],'url':doc.get('url',''),'doc_number':doc.get('doc_number',''),'effective_from':doc.get('effective_from',''),'verified':False,'text':content})
 (PROCESSED/'chunks.jsonl').write_text(''.join(json.dumps(x,ensure_ascii=False)+'\n' for x in chunks),encoding='utf-8')
 print('CHUNKS',len(chunks),'from',p)

def index():
 try:
  import chromadb
  from sentence_transformers import SentenceTransformer
 except ImportError:raise SystemExit('Install optional requirements-rag.txt first')
 p=PROCESSED/'chunks.jsonl'
 if not p.exists():raise SystemExit('First run process')
 arr=[json.loads(x) for x in p.read_text(encoding='utf-8').splitlines() if x.strip()]
 if not arr:print('No chunks');return
 model=SentenceTransformer('sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2')
 client=chromadb.PersistentClient(path=str(DATA/'chroma'));col=client.get_or_create_collection('gds_legal_v1',metadata={'hnsw:space':'cosine'})
 for n in range(0,len(arr),32):
  batch=arr[n:n+32];vectors=model.encode([x['text'] for x in batch],normalize_embeddings=True).tolist()
  col.upsert(ids=[x['id'] for x in batch],embeddings=vectors,documents=[x['text'] for x in batch],metadatas=[{'title':x['title'],'domain':x['domain'],'topic':x['topic'],'url':x['url'],'verified':False} for x in batch])
 print('VECTOR INDEX:',col.count(),'chunks')

if __name__=='__main__':
 a=argparse.ArgumentParser()
 a.add_argument('command',choices=['collect','process','index','import-local'])
 a.add_argument('--input-dir',help='Trusted administrator staging directory')
 a.add_argument('--manifest',help='CSV in input-dir; defaults to manifest.csv')
 args=a.parse_args()
 if args.command=='import-local':
  if not args.input_dir:a.error('import-local requires --input-dir')
  import_local(args.input_dir,args.manifest or str(Path(args.input_dir)/'manifest.csv'))
 else:{'collect':collect,'process':process,'index':index}[args.command]()
